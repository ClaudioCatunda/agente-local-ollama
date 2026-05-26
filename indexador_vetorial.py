import os
import httpx
import json
import chromadb
from pypdf import PdfReader
import openpyxl

def extrair_texto_pdf(caminho_arquivo):
    """Extrai o texto bruto de todas as páginas de um PDF."""
    try:
        leitor = PdfReader(caminho_arquivo)
        texto_completo = []
        for i, pagina in enumerate(leitor.pages):
            texto_pagina = pagina.extract_text()
            if texto_pagina:
                texto_completo.append(texto_pagina)
        return "\n".join(texto_completo)
    except Exception as e:
        print(f"   [Erro PDF]: Falha ao ler {os.path.basename(caminho_arquivo)}: {e}")
        return ""

def extrair_texto_excel(caminho_arquivo):
    """Converte as linhas e colunas de uma planilha em blocos de texto estruturado."""
    try:
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        conteudo_planilha = []
        
        for nome_aba in wb.sheetnames:
            aba = wb[nome_aba]
            conteudo_planilha.append(f"--- Aba: {nome_aba} ---")
            
            for linha in aba.iter_rows(values_only=True):
                if any(celula is not None for celula in linha):
                    linha_texto = " | ".join([str(celula) if celula is not None else "" for celula in linha])
                    conteudo_planilha.append(linha_texto)
                    
        return "\n".join(conteudo_planilha)
    except Exception as e:
        print(f"   [Erro Excel]: Falha ao ler {os.path.basename(caminho_arquivo)}: {e}")
        return ""
    
# Configurações de Caminhos no Catunda_SSD
RAIZ_SEGURA = "/Volumes/Catunda_SSD/Developer/Documents"
CHROMA_PATH = "/Volumes/Catunda_SSD/Developer/chroma_db"

# Configurações do Ollama
OLLAMA_EMBED_URL = "http://localhost:11434/api/embeddings"
MODELO_EMBED = "nomic-embed-text:latest"

def obter_embedding_ollama(texto: str) -> list:
    """Chama o modelo matemático local para gerar o vetor do fragmento de texto."""
    try:
        response = httpx.post(OLLAMA_EMBED_URL, json={
            "model": MODELO_EMBED,
            "prompt": texto
        }, timeout=30.0)
        
        if response.status_code == 200:
            return response.json().get("embedding")
        else:
            print(f"❌ Erro na API do Ollama: {response.status_code}")
    except Exception as e:
        print(f"❌ Falha de conexão ao gerar embedding: {e}")
    return []

def fatiar_texto(texto: str, tamanho_chunk: int = 1500, overlap: int = 200) -> list:
    """Corta o texto em blocos menores com uma zona de sobreposição (overlap) para não perder contexto."""
    chunks = []
    inicio = 0
    while inicio < len(texto):
        fim = inicio + tamanho_chunk
        chunks.append(texto[inicio:fim])
        inicio += (tamanho_chunk - overlap)
    return chunks

def executar_indexacao():
    print("📦 Inicializando Banco Vetorial ChromaDB no SSD...")
    cliente_chroma = chromadb.PersistentClient(path=CHROMA_PATH)
    colecao = cliente_chroma.get_or_create_collection(name="documentos_developer")
    
    # Suporte estendido para as novas extensões (Meta 1)
    extensoes_suportadas = ('.md', '.txt', '.pdf', '.xlsx')
    
    if not os.path.exists(RAIZ_SEGURA):
        print(f"❌ Diretório seguro não localizado em: {RAIZ_SEGURA}")
        return

    arquivos = [f for f in os.listdir(RAIZ_SEGURA) if os.path.isfile(os.path.join(RAIZ_SEGURA, f)) and f.endswith(extensoes_suportadas)]
    
    if not arquivos:
        print("⚠️ Nenhum arquivo válido encontrado para indexação.")
        return

    for arquivo in arquivos:
        # Ignora arquivos de controle interno e arquivos ocultos do sistema
        if arquivo in ["mapa_conhecimento.json", "INDICE_GERAL.md"] or arquivo.startswith('.'):
            continue
            
        caminho_completo = os.path.join(RAIZ_SEGURA, arquivo)
        extensao = os.path.splitext(arquivo)[1].lower()
        print(f"📂 Processando, fatiando e vetorizando: {arquivo}...")
        
        texto_bruto = ""
        try:
            # Roteamento baseado no formato do arquivo (Meta 1)
            if extensao in ('.md', '.txt'):
                with open(caminho_completo, 'r', encoding='utf-8', errors='ignore') as f:
                    texto_bruto = f.read()
            elif extensao == '.pdf':
                texto_bruto = extrair_texto_pdf(caminho_completo)
            elif extensao == '.xlsx':
                texto_bruto = extrair_texto_excel(caminho_completo)
                
            if not texto_bruto.strip():
                print(f"   ⚠️ Arquivo vazio ou sem texto extraível: {arquivo}")
                continue
            
            # Ajustado para 1500 caracteres com overlap de 200 para preservar o contexto das seções
            fragmentos = fatiar_texto(texto_bruto, tamanho_chunk=1500, overlap=200)
            
            contagem_salvos = 0
            for idx, fragmento in enumerate(fragmentos):
                if not fragmento.strip():
                    continue
                    
                doc_id = f"{arquivo}_chunk_{idx}"
                vetor = obter_embedding_ollama(fragmento)
                
                if vetor:
                    colecao.upsert(
                        ids=[doc_id],
                        embeddings=[vetor],
                        metadatas=[{"fonte": arquivo, "chunk": idx}],
                        documents=[fragmento]
                    )
                    contagem_salvos += 1
                    
            print(f"   ✅ {contagem_salvos} fragmentos salvos com sucesso para {arquivo}.")
            
        except Exception as e:
            print(f"❌ Erro ao processar o arquivo {arquivo}: {e}")

    print(f"\n🚀 Sucesso! Banco vetorial atualizado localmente em: {CHROMA_PATH}")

if __name__ == "__main__":
    executar_indexacao()